#!/usr/bin/env python3
"""
Script to read and display the contents of War Machine.xlsx
"""

import openpyxl
import sys


def read_excel_file(filename):
    """
    Read and display the contents of an Excel file.
    
    Args:
        filename (str): Path to the Excel file
    """
    try:
        # Load the workbook
        workbook = openpyxl.load_workbook(filename)
        
        # Print available sheets
        print(f"Excel file: {filename}")
        print(f"Available sheets: {workbook.sheetnames}")
        print("-" * 80)
        
        # Iterate through all sheets
        for sheet_name in workbook.sheetnames:
            worksheet = workbook[sheet_name]
            print(f"\nSheet: {sheet_name}")
            print(f"Dimensions: {worksheet.dimensions}")
            print("-" * 80)
            
            # Get all rows
            rows = list(worksheet.iter_rows(values_only=True))
            
            if not rows:
                print("(Empty sheet)")
                continue
            
            # Print header
            if rows:
                header = rows[0]
                print("Header:", header)
                print()
            
            # Print all data
            for i, row in enumerate(rows, start=1):
                print(f"Row {i}: {row}")
            
            print("-" * 80)
        
        workbook.close()
        print("\nSuccessfully read the Excel file!")
        return True
        
    except FileNotFoundError:
        print(f"Error: File '{filename}' not found.")
        return False
    except Exception as e:
        print(f"Error reading Excel file: {e}")
        return False


if __name__ == "__main__":
    # Default filename
    filename = "War Machine.xlsx"
    
    # Allow specifying a different file via command line
    if len(sys.argv) > 1:
        filename = sys.argv[1]
    
    success = read_excel_file(filename)
    sys.exit(0 if success else 1)
