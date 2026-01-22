#!/usr/bin/env python3
"""
Script to analyze and provide a comprehensive summary of the War Machine Excel file.
"""

import openpyxl
from collections import defaultdict


def analyze_war_machine():
    """Analyze the War Machine Excel file and provide a comprehensive summary."""
    
    filename = "War Machine.xlsx"
    workbook = openpyxl.load_workbook(filename)
    
    print("=" * 80)
    print("WAR MACHINE EXCEL FILE - COMPREHENSIVE ANALYSIS")
    print("=" * 80)
    print(f"\nFile: {filename}")
    print(f"Total Sheets: {len(workbook.sheetnames)}")
    print(f"Sheet Names: {', '.join(workbook.sheetnames)}")
    print("\n" + "=" * 80)
    
    # Analyze each sheet
    for sheet_name in workbook.sheetnames:
        worksheet = workbook[sheet_name]
        print(f"\n📊 SHEET: {sheet_name}")
        print("-" * 80)
        
        # Get all rows
        rows = list(worksheet.iter_rows(values_only=True))
        
        if not rows:
            print("  ⚠️  Empty sheet")
            continue
        
        # Find actual data (non-empty rows)
        non_empty_rows = [row for row in rows if any(cell is not None for cell in row)]
        
        print(f"  Dimensions: {worksheet.dimensions}")
        print(f"  Total Rows: {len(rows)}")
        print(f"  Non-Empty Rows: {len(non_empty_rows)}")
        
        if non_empty_rows:
            # Show header row
            header = non_empty_rows[0] if non_empty_rows else None
            if header:
                print(f"\n  📋 COLUMNS ({len([c for c in header if c])}):")
                for i, col in enumerate(header, 1):
                    if col:
                        print(f"    {i}. {col}")
            
            # Show first few data rows (skip header)
            if len(non_empty_rows) > 1:
                print(f"\n  📝 SAMPLE DATA (first 3 rows):")
                for i, row in enumerate(non_empty_rows[1:4], 1):
                    # Only show non-None values
                    data = [str(cell)[:30] if cell is not None else '' for cell in row]
                    non_empty_data = [(idx, val) for idx, val in enumerate(data) if val]
                    if non_empty_data:
                        print(f"    Row {i}:")
                        for idx, val in non_empty_data[:10]:  # Show first 10 columns
                            if header and idx < len(header) and header[idx]:
                                print(f"      {header[idx]}: {val}")
        
        print("-" * 80)
    
    workbook.close()
    
    print("\n" + "=" * 80)
    print("ANALYSIS COMPLETE")
    print("=" * 80)


if __name__ == "__main__":
    analyze_war_machine()
